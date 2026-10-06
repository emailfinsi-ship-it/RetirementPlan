"""Styling + native-chart helpers for the Excel export (see excel_export.py).

Everything visual about the generated workbook lives here so excel_export.py
can stay focused on *what* goes on each sheet: a consistent Arial type
scale, a dark header band with banded rows, money/percent formats, content
sized columns, frozen headers + filters, print setup, and openpyxl's native
(editable, live) Excel charts in the same categorical palette the dashboard
itself uses -- so a chart opened in Excel/Numbers/Sheets reads as the same
system as the pages on screen.
"""

from __future__ import annotations

from dataclasses import dataclass

from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT_NAME = "Arial"

INK = "1D1D1F"
MUTED = "6E6E73"
BRAND = "0071E3"
BAND = "F5F5F7"
HAIRLINE = "D2D2D7"
KEY_FILL = "E8F1FC"  # the dashboard's own --brand-soft

# Same fixed-order categorical slots as the dashboard's donut charts.
PALETTE = ("2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300", "898781")

MONEY0 = '$#,##0;($#,##0);"-"'
MONEY2 = '$#,##0.00;($#,##0.00);"-"'
PCT1 = "0.0%"
PCT2 = "0.00%"
INT = "#,##0"
YEAR = "0"

_THIN = Side(style="thin", color=HAIRLINE)
_HEADER_FILL = PatternFill("solid", fgColor=INK)
_BAND_FILL = PatternFill("solid", fgColor=BAND)
_KEY_FILL = PatternFill("solid", fgColor=KEY_FILL)


@dataclass
class Fmt:
    """A cell value that carries its own number format -- for the
    label/value tables (Summary, Monte Carlo metrics) where a single
    column mixes dollars, percentages and plain counts."""

    value: object
    fmt: str


@dataclass
class TableInfo:
    header_row: int
    first_row: int  # first data row (0 data rows => first_row > last_row)
    last_row: int  # last data row, INCLUDING any trailing total rows
    first_col: int
    last_col: int
    next_row: int  # first free row after the table


def font(*, size: int = 10, bold: bool = False, color: str = INK, italic: bool = False) -> Font:
    return Font(name=FONT_NAME, size=size, bold=bold, color=color, italic=italic)


def write_title(ws, title: str, subtitle: str = "") -> int:
    """Big sheet title (+ muted subtitle) in A1/A2; returns the first free row."""
    ws["A1"] = title
    ws["A1"].font = font(size=18, bold=True)
    ws.row_dimensions[1].height = 28
    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = font(size=10, color=MUTED)
    return 4


def write_section_header(ws, row: int, col: int, text: str, *, span: int = 4) -> int:
    """A section heading with a brand-blue rule under it; returns the next row (where the table goes)."""
    cell = ws.cell(row=row, column=col, value=text)
    cell.font = font(size=12, bold=True)
    for c in range(col, col + span):
        ws.cell(row=row, column=c).border = Border(bottom=Side(style="medium", color=BRAND))
    ws.row_dimensions[row].height = 20
    return row + 1


def write_table(
    ws,
    *,
    headers: list,
    rows: list,
    money_cols: set = frozenset(),
    pct_cols: set = frozenset(),
    money_format: str = MONEY0,
    pct_format: str = PCT2,
    year_cols: set = frozenset(),
    int_cols: set = frozenset(),
    wrap_cols: set = frozenset(),
    start_row: int = 1,
    start_col: int = 1,
    total_rows: int = 0,
    key_rows: set = frozenset(),
    freeze: bool = False,
    autofilter: bool = False,
    header_height: float | None = 24,
) -> TableInfo:
    """Writes a styled header row + data rows. `*_cols` are 0-based offsets
    from start_col. `total_rows` trailing rows are drawn as totals (bold,
    top rule, excluded from the autofilter); `key_rows` (0-based data-row
    offsets) get the highlight fill. Cell values may be wrapped in Fmt() to
    override the column format for that one cell."""
    for c_offset, header in enumerate(headers):
        cell = ws.cell(row=start_row, column=start_col + c_offset, value=header)
        cell.font = font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        numeric = c_offset in money_cols or c_offset in pct_cols or c_offset in int_cols
        cell.alignment = Alignment(
            horizontal="right" if numeric else "left", vertical="center", wrap_text=True,
            indent=0 if numeric else 1,
        )
    if header_height:
        ws.row_dimensions[start_row].height = header_height

    last_row = start_row
    n_rows = len(rows)
    for r_offset, row in enumerate(rows):
        r = start_row + 1 + r_offset
        last_row = r
        is_total = total_rows and r_offset >= n_rows - total_rows
        for c_offset, value in enumerate(row):
            cell = ws.cell(row=r, column=start_col + c_offset)
            fmt = None
            if isinstance(value, Fmt):
                value, fmt = value.value, value.fmt
            cell.value = value
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                if fmt is None:
                    if c_offset in money_cols:
                        fmt = money_format
                    elif c_offset in pct_cols:
                        fmt = pct_format
                    elif c_offset in year_cols:
                        fmt = YEAR
                    elif c_offset in int_cols:
                        fmt = INT
                if fmt:
                    cell.number_format = fmt
                cell.alignment = Alignment(
                    horizontal="left" if (c_offset in year_cols and fmt is None) else "right", vertical="top",
                    indent=1 if c_offset in year_cols else 0,
                )
            else:
                cell.alignment = Alignment(
                    horizontal="left", vertical="top", wrap_text=c_offset in wrap_cols, indent=1,
                )
            cell.font = font(bold=bool(is_total) or r_offset in key_rows)
            if is_total:
                cell.border = Border(top=Side(style="medium", color=INK), bottom=_THIN)
            else:
                cell.border = Border(bottom=_THIN)
                if r_offset in key_rows:
                    cell.fill = _KEY_FILL
                elif r_offset % 2 == 1:
                    cell.fill = _BAND_FILL

    last_col = start_col + max(len(headers), 1) - 1
    if freeze and start_row == 1 and start_col == 1:
        ws.freeze_panes = ws.cell(row=2, column=1)
    if autofilter and rows:
        filter_last = last_row - total_rows if total_rows else last_row
        if filter_last > start_row:
            ws.auto_filter.ref = (
                f"{get_column_letter(start_col)}{start_row}:{get_column_letter(last_col)}{filter_last}"
            )
    return TableInfo(start_row, start_row + 1, last_row, start_col, last_col, last_row + 1)


def autosize(ws, *, skip_rows: frozenset | set = frozenset(), min_width: int = 10, max_width: int = 46) -> None:
    """Content-sized columns (header or widest value, clamped), skipping
    title rows so a long banner never stretches column A."""
    widths: dict[int, float] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None or cell.row in skip_rows or type(cell).__name__ == "MergedCell":
                continue
            value = cell.value
            if isinstance(value, float):
                text = f"{value:,.2f}"
                length = len(text) + 2
            elif isinstance(value, int):
                length = len(f"{value:,}") + 2
            else:
                length = max((len(line) for line in str(value).split("\n")), default=0)
            if cell.font is not None and cell.font.bold:
                length = length * 1.15 + 1
            widths[cell.column] = max(widths.get(cell.column, 0), length)
    for col, length in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(min_width, min(max_width, length + 3))


def finish_sheet(ws, *, tab_color: str, gridlines: bool = True, landscape: bool = True, one_page: bool = False) -> None:
    ws.sheet_properties.tabColor = tab_color
    ws.sheet_view.showGridLines = gridlines
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1 if one_page else 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = False
    ws.page_margins.left = ws.page_margins.right = 0.5
    ws.oddFooter.center.text = "Retirement Plan -- &A -- page &P of &N"


# --------------------------------------------------------------------------
# Charts
# --------------------------------------------------------------------------


def _line_props(color: str, *, width_pt: float = 2.25, dash: str | None = None) -> GraphicalProperties:
    props = GraphicalProperties()
    props.line = LineProperties(solidFill=color, w=int(width_pt * 12700), prstDash=dash)
    props.noFill = True
    return props


def _style_axes(chart, *, y_format: str | None, x_title: str | None = None, y_title: str | None = None) -> None:
    for axis in (chart.x_axis, chart.y_axis):
        axis.delete = False
    chart.x_axis.tickLblPos = "low"  # category labels stay under the plot even when values go negative
    chart.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E5E5EA", w=9525))
    chart.y_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
    chart.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill=HAIRLINE, w=9525))
    if y_format:
        chart.y_axis.number_format = y_format
        chart.y_axis.numFmt.sourceLinked = False
    if x_title:
        chart.x_axis.title = x_title
    if y_title:
        chart.y_axis.title = y_title


def _labels(*, show_val: bool = False, show_percent: bool = False, num_fmt: str | None = None) -> DataLabelList:
    """Data labels with every other element (series/category name, legend key) explicitly off --
    left unset, LibreOffice/Excel differ on what to show."""
    labels = DataLabelList()
    labels.showVal = show_val
    labels.showPercent = show_percent
    labels.showCatName = False
    labels.showSerName = False
    labels.showLegendKey = False
    labels.showLeaderLines = False
    if num_fmt:
        labels.numFmt = num_fmt
    return labels


def _base(chart, title: str, width: float, height: float):
    chart.title = title
    chart.width, chart.height = width, height
    chart.legend.position = "b"
    return chart


def add_line_chart(
    ws, *, title: str, anchor: str, cats: tuple, series: list, y_format: str | None = '$#,##0.0,,"M"',
    width: float = 20, height: float = 9.5, x_title: str | None = None, y_title: str | None = None,
    y_min: float | None = None,
) -> None:
    """series: [(title, (min_col, min_row, max_row), color, dash|None, width_pt)] -- references into `ws`.
    cats: (col, min_row, max_row) for the x-axis labels."""
    chart = _base(LineChart(), title, width, height)
    _style_axes(chart, y_format=y_format, x_title=x_title, y_title=y_title)
    if y_min is not None:
        chart.y_axis.scaling.min = y_min
    for s_title, (col, r1, r2), color, dash, pt in series:
        s = Series(Reference(ws, min_col=col, min_row=r1, max_row=r2), title=s_title)
        s.graphicalProperties = _line_props(color, width_pt=pt, dash=dash)
        s.marker.symbol = "none"
        s.smooth = False
        chart.series.append(s)
    col, r1, r2 = cats
    chart.set_categories(Reference(ws, min_col=col, min_row=r1, max_row=r2))
    ws.add_chart(chart, anchor)


def add_bar_chart(
    ws, *, title: str, anchor: str, cats: tuple, series: list, horizontal: bool = False, stacked: bool = False,
    y_format: str | None = MONEY0, width: float = 16, height: float = 9, legend: bool = True,
    data_labels: bool = False,
) -> None:
    """series: [(title, (col, min_row, max_row), color)]."""
    chart = _base(BarChart(), title, width, height)
    chart.type = "bar" if horizontal else "col"
    if stacked:
        chart.grouping = "stacked"
        chart.overlap = 100
    chart.gapWidth = 60
    _style_axes(chart, y_format=y_format)
    for s_title, (col, r1, r2), color in series:
        s = Series(Reference(ws, min_col=col, min_row=r1, max_row=r2), title=s_title)
        s.graphicalProperties = GraphicalProperties(solidFill=color)
        s.graphicalProperties.line = LineProperties(noFill=True)
        chart.series.append(s)
    col, r1, r2 = cats
    chart.set_categories(Reference(ws, min_col=col, min_row=r1, max_row=r2))
    if horizontal:
        chart.x_axis.scaling.orientation = "maxMin"  # first row on top
        chart.y_axis.crosses = "max"  # ...and keep the value axis along the bottom
    if not legend or len(series) == 1:
        chart.legend = None
    if data_labels:
        chart.dataLabels = _labels(show_val=True, num_fmt=y_format or MONEY0)
    ws.add_chart(chart, anchor)


def add_doughnut_chart(
    ws, *, title: str, anchor: str, cats: tuple, values: tuple, width: float = 13, height: float = 9.5,
) -> None:
    """One slice per category row, in fixed palette order (never cycled past
    the 7th slot -- callers fold the tail into "Other" first)."""
    chart = DoughnutChart()
    chart.title = title
    chart.width, chart.height = width, height
    chart.holeSize = 55
    chart.legend.position = "r"
    col, r1, r2 = values
    s = Series(Reference(ws, min_col=col, min_row=r1, max_row=r2), title=title)
    for i in range(r2 - r1 + 1):
        point = DataPoint(idx=i)
        point.graphicalProperties = GraphicalProperties(solidFill=PALETTE[min(i, len(PALETTE) - 1)])
        point.graphicalProperties.line = LineProperties(solidFill="FFFFFF", w=19050)
        s.dPt.append(point)
    s.dLbls = _labels(show_percent=True)
    chart.series.append(s)
    ccol, cr1, cr2 = cats
    chart.set_categories(Reference(ws, min_col=ccol, min_row=cr1, max_row=cr2))
    ws.add_chart(chart, anchor)


def top_n_with_other(items: list[tuple[str, float]], n: int = 6) -> list[tuple[str, float]]:
    """Largest-first (positive values only), tail folded into "Other" -- the
    same rule the dashboard's donut charts apply."""
    positive = sorted(((k, v) for k, v in items if v and v > 0), key=lambda kv: kv[1], reverse=True)
    if len(positive) <= n + 1:
        return positive
    head, tail = positive[:n], positive[n:]
    return head + [("Other", sum(v for _, v in tail))]


def fitted_axis_min(values: list[float]) -> float | None:
    """A y-axis floor a little under the smallest value (rounded down to a
    tidy step) so a line that only moves a few percent isn't flattened
    against a $0 baseline -- same idea as the dashboard's net worth chart."""
    values = [v for v in values if v is not None]
    if not values:
        return None
    lo, hi = min(values), max(values)
    if lo <= 0 or hi <= 0:
        return None
    import math

    step = 10 ** max(0, int(math.floor(math.log10(hi))) - 1)
    return math.floor(lo * 0.95 / step) * step
