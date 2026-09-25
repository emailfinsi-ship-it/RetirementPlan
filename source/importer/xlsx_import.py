"""One-time import of the user's retirement-plan workbook into this
program's own database.

This is the ONLY place this program ever reads that spreadsheet. After a
successful import, this program is the system of record: line items are
added, edited, and deleted here (source.common.db), and a fresh .xlsx can
be generated any time going forward (source.reporting.excel_export) --
there is no ongoing dependency on the original file, and no attempt to
keep re-reading it in sync.

Matches the user's real workbook (e.g. "Retirement_Plan_2028-2060.xlsx"):
sheets "Summary & Assumptions" (401K/IRA and Liquid Investments account
blocks, Loans/Committed, Annual Expense Need, Retirement Income
Assumptions, Key Modeling Assumptions) and "Real Estate" (one row per
property). "Yearly Drawdown", "Monte Carlo Simulation", "MC - S&P500
Historical", "Historical Stress Test", and "MC - 60-40 Blend" are all
*computed* output sheets from the old workbook's own model -- this
program computes those itself (source.reporting.monte_carlo), so they're
never read as import input, with one small exception: "Monte Carlo
Simulation" carries two assumptions (return volatility, trial count)
that aren't written anywhere on Summary & Assumptions, so they're read
from there too, best-effort, if that sheet happens to be present.

Earlier versions of this program (and the retired FamilyFinanceTracker
retirement-plan feature before it) were built against a differently
laid-out workbook (Summary / Asset / Expense / RealEstate Analysis /
Monte Carlo sheets, fixed fixed row/column coordinates). That layout is
no longer supported -- this reader is built for the current one.

Rows are located by searching each sheet's first column for a label
(case-insensitive prefix match, e.g. "annual rental income" matches the
cell reading "Annual rental income (net)") rather than trusting a fixed
row number, so inserting or removing a row elsewhere on the sheet (an
extra 401k account, say) doesn't silently misalign every reader below
it. Re-running the import (e.g. from the Setup page, "Re-import") clears
out only previously-imported rows (source='imported') from each section
first, so anything the user has since typed in by hand (source='manual')
is left untouched.
"""

from __future__ import annotations

import io
import re
import time
from pathlib import Path

import msoffcrypto
import openpyxl

from source.common import db

EXPECTED_SHEETS = ("Summary & Assumptions", "Real Estate")


class ImportError_(RuntimeError):
    """Workbook missing, unreadable, wrong password, or a sheet this
    importer expects isn't there (moved/renamed/deleted)."""


# --------------------------------------------------------------------
# Decrypt + load (the plaintext workbook bytes only ever exist in
# memory, never written to disk, and the password is never logged or
# echoed back in an error).
# --------------------------------------------------------------------


def _decrypt_to_buffer(path: Path, password: str | None) -> io.BytesIO:
    buf = io.BytesIO()
    with path.open("rb") as fh:
        office_file = msoffcrypto.OfficeFile(fh)
        if office_file.is_encrypted():
            if not password:
                raise ImportError_(
                    "This workbook is password-protected. Enter its password on the Setup page and try again."
                )
            try:
                office_file.load_key(password=password)
                office_file.decrypt(buf)
            except Exception as exc:  # msoffcrypto raises plain Exception on a bad password
                raise ImportError_(
                    "Could not open the workbook -- the password entered is probably wrong."
                ) from exc
        else:
            fh.seek(0)
            buf.write(fh.read())
    buf.seek(0)
    return buf


def load_workbook(path: Path, *, password: str | None = None):
    if not path.exists():
        raise ImportError_(f"No workbook found at {path}.")
    buf = _decrypt_to_buffer(path, password)
    try:
        wb = openpyxl.load_workbook(buf, data_only=True, read_only=True)
    except Exception as exc:
        raise ImportError_(
            "The workbook opened but isn't a readable .xlsx -- was it actually exported as Excel?"
        ) from exc
    missing = [name for name in EXPECTED_SHEETS if name not in wb.sheetnames]
    if missing:
        raise ImportError_(
            "This workbook is missing expected sheet(s): " + ", ".join(missing) + ". "
            f"Sheets found: {', '.join(wb.sheetnames)}."
        )
    return wb


# --------------------------------------------------------------------
# Small generic cell helpers
# --------------------------------------------------------------------


def _cell(ws, row: int, col: int):
    return ws.cell(row=row, column=col).value


def _num(value, default: float = 0.0) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _ownership_pct(raw) -> float:
    """Stored either as a fraction (0.5) or an already-multiplied
    percentage (50/100) depending on the cell's own number format,
    which openpyxl's data_only read can't tell apart -- so treat
    anything at or below 1 as a fraction and scale it up, and anything
    above 1 as already a percentage. A blank cell defaults to full
    (100%) ownership, the common case for a single-owner property."""
    if raw is None:
        return 100.0
    value = _num(raw, default=1.0)
    return value * 100.0 if value <= 1 else value


def _find_label_row(ws, prefix: str, *, col: int = 1) -> int | None:
    """Row number of the first cell in `col` (top to bottom) whose text,
    lowercased and stripped, starts with `prefix` -- None if no row
    matches. Case/prefix matching (rather than exact equality) survives
    the workbook's own punctuation ("Plan horizon — end age" vs a
    shorter search prefix "plan horizon") without needing to reproduce
    an em dash or exact wording in this code."""
    target = prefix.strip().lower()
    for r in range(1, ws.max_row + 1):
        value = _cell(ws, r, col)
        if value and str(value).strip().lower().startswith(target):
            return r
    return None


def _find_value(ws, prefix: str, *, label_col: int = 1, value_col: int = 2):
    """The value in `value_col` on the row whose `label_col` cell starts
    with `prefix` -- None if that label isn't found on this sheet."""
    row = _find_label_row(ws, prefix, col=label_col)
    return _cell(ws, row, value_col) if row is not None else None


_BASIS_YEAR_RE = re.compile(r"\((\d{4})\s*basis\)", re.IGNORECASE)


def _parse_basis_year(label, default: int = 2026) -> int:
    """Pulls the year out of a label like "Annual Expense Need (2026
    basis)" -- falls back to `default` if the label is missing or
    doesn't have that "(YYYY basis)" phrasing."""
    if not label:
        return default
    m = _BASIS_YEAR_RE.search(str(label))
    return int(m.group(1)) if m else default


# --------------------------------------------------------------------
# Per-sheet readers
# --------------------------------------------------------------------


def _read_real_estate_rows(wb) -> list[dict]:
    """Real Estate sheet: one data row per property (name, purchase
    price, current value, ownership %, loan balance), stopping at the
    sheet's own Total row or the end of its used range."""
    ws = wb["Real Estate"]
    header_row = _find_label_row(ws, "asset") or 3
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        name = _cell(ws, r, 1)
        if not name:
            continue
        label = str(name).strip()
        if label.lower().startswith("total"):
            break
        current_value = _cell(ws, r, 3)
        if not isinstance(current_value, (int, float)):
            continue
        rows.append({
            "name": label,
            "purchase_price": _num(_cell(ws, r, 2)),
            "current_value": _num(current_value),
            "ownership_pct": _ownership_pct(_cell(ws, r, 4)),
            "loan_balance": _num(_cell(ws, r, 6)),
        })
    return rows


def _read_labeled_balance_block(ws, label_prefix: str, *, name_col: int = 1, balance_col: int = 2) -> list[dict]:
    """Reads the two-column name/balance block that starts two rows
    below a section-header cell whose text starts with `label_prefix`
    (the header row itself, then an "Account"/"Balance" or
    "Item"/"Balance" sub-header row, then data rows) -- stops at the
    first blank name or a name starting with "total" (that block's own
    total row). Returns [] if the label isn't found on this sheet (a
    future export renamed or removed a section) rather than raising, so
    one missing block doesn't fail the whole import."""
    header_row = _find_label_row(ws, label_prefix, col=name_col)
    if header_row is None:
        return []
    rows = []
    r = header_row + 2
    while True:
        name = _cell(ws, r, name_col)
        if not name:
            break
        label = str(name).strip()
        if label.lower().startswith("total"):
            break
        balance = _cell(ws, r, balance_col)
        if isinstance(balance, (int, float)):
            rows.append({"name": label, "balance": _num(balance)})
        r += 1
    return rows


def _read_extra_mc_assumptions(wb) -> dict:
    """Best-effort extras (return volatility, trial count) from the
    Monte Carlo Simulation sheet, if present -- not in EXPECTED_SHEETS,
    since only Summary & Assumptions and Real Estate are load-bearing
    for the core Asset/Real Estate/Expense/Income data; every other
    sheet in the workbook (including this one) is the old model's own
    *computed* output, superseded by this program's own Monte Carlo
    engine, except for these two input assumptions which aren't
    recorded anywhere on Summary & Assumptions."""
    if "Monte Carlo Simulation" not in wb.sheetnames:
        return {}
    ws = wb["Monte Carlo Simulation"]
    result = {}
    volatility = _find_value(ws, "return volatility")
    if isinstance(volatility, (int, float)):
        result["return_volatility"] = volatility
    trials = _find_value(ws, "number of trials")
    if isinstance(trials, (int, float)):
        result["iterations"] = int(trials)
    return result


# --------------------------------------------------------------------
# Top-level import
# --------------------------------------------------------------------


def import_from_xlsx(path: Path, *, password: str | None = None) -> dict:
    """Reads the workbook and populates real_estate / assets / expenses /
    income / assumptions. Clears out previously-imported rows
    (source='imported') in each section first, so re-running this after
    fixing a mistake in the source workbook doesn't just keep piling up
    duplicates -- but never touches rows the user added by hand
    (source='manual'). Returns a summary dict of what was imported."""
    wb = load_workbook(path, password=password)
    sw = wb["Summary & Assumptions"]

    real_estate_rows = _read_real_estate_rows(wb)
    k401_rows = _read_labeled_balance_block(sw, "401k / ira")
    liquid_rows = _read_labeled_balance_block(sw, "liquid investments")
    loan_rows = _read_labeled_balance_block(sw, "loans / committed")

    basis_row = _find_label_row(sw, "annual expense need")
    basis_year = _parse_basis_year(_cell(sw, basis_row, 1) if basis_row is not None else None)
    fixed_cost = _num(_find_value(sw, "fixed cost"))
    variable_cost = _num(_find_value(sw, "variable cost"))

    rental_income = _num(_find_value(sw, "annual rental income"))
    ss_annual = _num(_find_value(sw, "social security (annual)"))
    ss_start_age = _find_value(sw, "ss start age")
    birth_year = _find_value(sw, "birth year")
    retirement_year = _find_value(sw, "retirement year")
    retirement_age = _find_value(sw, "retirement age")
    growth_rate = _find_value(sw, "portfolio growth rate")
    inflation_rate = _find_value(sw, "inflation rate")
    plan_horizon_end_age = _find_value(sw, "plan horizon")

    extra_mc_assumptions = _read_extra_mc_assumptions(wb)

    db.clear_section("real_estate", source="imported")
    db.clear_section("assets", source="imported")
    db.clear_section("expenses", source="imported")
    db.clear_section("income", source="imported")

    for row in real_estate_rows:
        db.add_line_item("real_estate", {**row, "source": "imported"})

    asset_count = 0
    for category, rows in (("401k", k401_rows), ("Liquid", liquid_rows)):
        for row in rows:
            db.add_line_item("assets", {
                "kind": "asset",
                "category": category,
                "name": row["name"],
                "balance": row["balance"],
                "source": "imported",
            })
            asset_count += 1
    for row in loan_rows:
        db.add_line_item("assets", {
            "kind": "liability",
            "category": "Loan",
            "name": row["name"],
            "balance": row["balance"],
            "source": "imported",
        })
        asset_count += 1

    expense_count = 0
    if fixed_cost:
        db.add_line_item("expenses", {
            "name": "Fixed cost (house + property tax)",
            "category": "Fixed",
            "annual_amount": fixed_cost,
            "base_year": basis_year,
            "inflates": 1,
            "source": "imported",
        })
        expense_count += 1
    if variable_cost:
        db.add_line_item("expenses", {
            "name": "Variable cost (utility, medical, food, etc.)",
            "category": "Variable",
            "annual_amount": variable_cost,
            "base_year": basis_year,
            "inflates": 1,
            "source": "imported",
        })
        expense_count += 1

    income_count = 0
    if rental_income:
        db.add_line_item("income", {
            "name": "Rental income",
            "category": "Rental",
            "annual_amount": rental_income,
            "base_year": basis_year,
            "inflates": 1,
            "source": "imported",
        })
        income_count += 1
    if ss_annual:
        ss_start_year = None
        if isinstance(birth_year, (int, float)) and isinstance(ss_start_age, (int, float)):
            ss_start_year = int(birth_year) + int(ss_start_age)
        db.add_line_item("income", {
            "name": "Social Security",
            "category": "Social Security",
            "annual_amount": ss_annual,
            "start_year": ss_start_year,
            "inflates": 0,
            "source": "imported",
        })
        income_count += 1

    assumption_updates = {}
    if isinstance(growth_rate, (int, float)):
        assumption_updates["expected_return"] = growth_rate
    if isinstance(inflation_rate, (int, float)):
        assumption_updates["inflation_rate"] = inflation_rate
    if isinstance(retirement_year, (int, float)):
        assumption_updates["current_year"] = int(retirement_year)
    if isinstance(retirement_age, (int, float)) and isinstance(plan_horizon_end_age, (int, float)):
        assumption_updates["plan_horizon_years"] = int(plan_horizon_end_age) - int(retirement_age)
    assumption_updates.update(extra_mc_assumptions)
    if assumption_updates:
        db.set_assumptions(assumption_updates)

    db.set_meta("last_import_at", time.strftime("%Y-%m-%d %H:%M:%S"))
    db.set_meta("last_import_source", str(path))

    return {
        "real_estate_count": len(real_estate_rows),
        "asset_count": asset_count,
        "expense_count": expense_count,
        "income_count": income_count,
        "assumptions_imported": sorted(assumption_updates.keys()),
    }
