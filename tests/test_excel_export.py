import unittest

import openpyxl

from source.common import db
from source.reporting import excel_export
from tests.testing_utils import TempConfigTestCase


class TestExcelExport(TempConfigTestCase):
    def setUp(self):
        super().setUp()
        db.add_line_item("assets", {"kind": "asset", "category": "401k", "name": "Fidelity", "balance": 100_000, "include_in_monte_carlo": 1})
        db.add_line_item("expenses", {"name": "Baseline", "annual_amount": 40_000})
        db.add_line_item("real_estate", {"name": "Home", "current_value": 500_000, "loan_balance": 200_000})

    def test_export_writes_a_readable_workbook_with_expected_sheets(self):
        path = excel_export.export_current_state()
        self.assertTrue(path.exists())
        wb = openpyxl.load_workbook(str(path))
        self.assertEqual(
            set(wb.sheetnames),
            {
                "Summary", "Asset", "Expense", "Income", "Real Estate Analysis",
                "Linked Positions", "Linked Accounts", "Monte Carlo Simulation", "Monte Carlo Stress Test",
                "Net Worth History",
            },
        )

    def test_asset_sheet_contains_the_line_item(self):
        path = excel_export.export_current_state()
        wb = openpyxl.load_workbook(str(path))
        ws = wb["Asset"]
        names = [ws.cell(row=r, column=3).value for r in range(2, ws.max_row + 1)]
        self.assertIn("Fidelity", names)

    def test_export_records_meta(self):
        excel_export.export_current_state()
        self.assertIsNotNone(db.get_meta("last_export_at"))
        self.assertIsNotNone(db.get_meta("last_export_path"))

    def test_real_estate_sheet_has_annual_income_column(self):
        path = excel_export.export_current_state()
        wb = openpyxl.load_workbook(str(path))
        ws = wb["Real Estate Analysis"]
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        self.assertIn("Annual Income", headers)

    def test_real_estate_sheet_your_loan_share_reflects_ownership_pct(self):
        db.clear_section("real_estate")
        db.add_line_item("real_estate", {
            "name": "Moorpark", "current_value": 900_000, "loan_balance": 400_000, "ownership_pct": 50,
        })
        path = excel_export.export_current_state()
        wb = openpyxl.load_workbook(str(path))
        ws = wb["Real Estate Analysis"]
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        share_col = headers.index("Your Loan Share") + 1
        equity_col = headers.index("Equity") + 1
        row = next(r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value == "Moorpark")
        self.assertEqual(ws.cell(row=row, column=share_col).value, 200_000)
        self.assertEqual(ws.cell(row=row, column=equity_col).value, 250_000)

    def test_asset_sheet_exports_live_balance_for_a_linked_row(self):
        db.upsert_plaid_item(
            item_key="fidelity", plaid_item_id="item-1", institution_id="ins_1",
            institution_name="Fidelity", products=["investments"], environment="sandbox",
            linked_at="2026-09-01T00:00:00+00:00",
        )
        db.upsert_plaid_account(item_key="fidelity", account={
            "account_id": "acc-1", "name": "401k", "official_name": "Fidelity 401k",
            "type": "investment", "subtype": "401k", "mask": "1234",
            "balances": {"current": 55_000.0, "available": None, "iso_currency_code": "USD"},
        })
        item_id = db.add_line_item("assets", {"kind": "asset", "name": "Linked 401k", "balance": 1})
        db.update_line_item("assets", item_id, {"linked_account_id": "acc-1"})
        path = excel_export.export_current_state()
        wb = openpyxl.load_workbook(str(path))
        ws = wb["Asset"]
        rows = {ws.cell(row=r, column=3).value: ws.cell(row=r, column=6).value for r in range(2, ws.max_row + 1)}
        self.assertEqual(rows["Linked 401k"], 55_000.0)

    def test_monte_carlo_band_written_when_result_supplied(self):
        result = {"band": [{"year": 2027, "p10": 1, "median": 2, "p90": 3}]}
        path = excel_export.export_current_state(monte_carlo_result=result)
        wb = openpyxl.load_workbook(str(path))
        ws = wb["Monte Carlo Simulation"]
        values = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
        self.assertIn(2027, values)

    def test_monte_carlo_sheet_self_computes_when_no_result_is_supplied(self):
        # Regression test: /api/export used to call export_current_state()
        # with no monte_carlo_result at all, so the Monte Carlo Simulation
        # sheet's band table silently never had any data in it. It should
        # now run the simulations itself.
        path = excel_export.export_current_state()
        wb = openpyxl.load_workbook(str(path))
        ws = wb["Monte Carlo Simulation"]
        values = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
        self.assertIn("Gaussian simulation (Expected return / Volatility)", values)
        self.assertIn("S&P 500 historical-returns simulation", values)

    def test_summary_sheet_includes_net_liquid_asset_and_category_subtotals(self):
        path = excel_export.export_current_state()
        wb = openpyxl.load_workbook(str(path))
        ws = wb["Summary"]
        labels = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
        self.assertIn("Net Liquid Asset (Monte Carlo starting balance)", labels)
        self.assertIn("401k", labels)  # asset category subtotal row

    def test_linked_positions_sheet_contains_synced_holdings(self):
        db.upsert_plaid_item(
            item_key="fidelity", plaid_item_id="item-1", institution_id="ins_1",
            institution_name="Fidelity", products=["investments"], environment="sandbox",
            linked_at="2026-09-01T00:00:00+00:00",
        )
        db.upsert_plaid_account(item_key="fidelity", account={
            "account_id": "acc-1", "name": "Brokerage", "official_name": "Fidelity Brokerage",
            "type": "investment", "subtype": "brokerage", "mask": "1234",
            "balances": {"current": 10_000.0, "available": None, "iso_currency_code": "USD"},
        })
        db.upsert_plaid_holding(item_key="fidelity", holding={
            "account_id": "acc-1", "security_id": "sec-1", "ticker_symbol": "AAPL",
            "security_name": "Apple Inc.", "security_type": "equity", "quantity": 10,
            "institution_price": 200.0, "institution_value": 2_000.0, "cost_basis": 1_500.0,
        })
        path = excel_export.export_current_state()
        wb = openpyxl.load_workbook(str(path))
        ws = wb["Linked Positions"]
        names = [ws.cell(row=r, column=3).value for r in range(2, ws.max_row + 1)]
        self.assertIn("Apple Inc.", names)

    def test_linked_accounts_sheet_contains_synced_account_balance(self):
        db.upsert_plaid_item(
            item_key="chase", plaid_item_id="item-2", institution_id="ins_2",
            institution_name="Chase", products=["auth"], environment="sandbox",
            linked_at="2026-09-01T00:00:00+00:00",
        )
        db.upsert_plaid_account(item_key="chase", account={
            "account_id": "acc-2", "name": "Checking", "official_name": "Chase Checking",
            "type": "depository", "subtype": "checking", "mask": "5678",
            "balances": {"current": 3_000.0, "available": 3_000.0, "iso_currency_code": "USD"},
        })
        path = excel_export.export_current_state()
        wb = openpyxl.load_workbook(str(path))
        ws = wb["Linked Accounts"]
        # Column 1 is the real institution name (the parent Item's
        # institution_name, "Chase" here) -- not the account's own Plaid
        # name ("Checking"), which this sheet doesn't otherwise surface.
        balances = {
            ws.cell(row=r, column=1).value: ws.cell(row=r, column=4).value for r in range(2, ws.max_row + 1)
        }
        self.assertEqual(balances["Chase"], 3_000.0)

    def test_stress_test_sheet_has_expense_income_drawdown_and_scenario_columns(self):
        # export_current_state() with no monte_carlo_result self-computes
        # via _default_monte_carlo_result(), which passes pools= the same
        # way the live Monte Carlo tab does -- so this sheet is always in
        # per-scenario-tax mode (retirement_pools is never None), not the
        # legacy shared-columns layout.
        path = excel_export.export_current_state()
        wb = openpyxl.load_workbook(str(path))
        ws = wb["Monte Carlo Stress Test"]
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        self.assertEqual(headers[:4], ["Year", "Total Expense", "Total Income", "Yearly Drawdown"])
        self.assertNotIn("Tax (Fed+CA)", headers)
        self.assertIn("Great Depression (1929) Return", headers)
        self.assertIn("Great Depression (1929) Tax", headers)
        self.assertIn("Great Depression (1929) Withdrawal", headers)
        self.assertIn("Great Depression (1929) Balance", headers)

    def test_stress_test_sheet_tax_column_reflects_the_apply_drawdown_tax_assumption(self):
        # Drawdown tax is on by default (source.common.db.DEFAULT_
        # ASSUMPTIONS) -- with a real withdrawal need, the first
        # scenario's first-year Tax cell should be a positive number and
        # Withdrawal should be Yearly Drawdown + Tax, not just Yearly
        # Drawdown. (Per-scenario mode -- see the test above -- so this
        # reads the first scenario's own Tax/Withdrawal columns, not a
        # shared pair.)
        db.add_line_item("assets", {"kind": "asset", "category": "401k", "name": "Fidelity 401k", "balance": 1_000_000, "include_in_monte_carlo": 1})
        db.add_line_item("expenses", {"name": "Baseline", "annual_amount": 100_000, "start_year": None, "end_year": None})
        try:
            path = excel_export.export_current_state()
            wb = openpyxl.load_workbook(str(path))
            ws = wb["Monte Carlo Stress Test"]
            headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
            tax_col = headers.index("Great Depression (1929) Tax") + 1
            withdrawal_col = headers.index("Great Depression (1929) Withdrawal") + 1
            drawdown_col = headers.index("Yearly Drawdown") + 1
            tax = ws.cell(row=2, column=tax_col).value
            drawdown = ws.cell(row=2, column=drawdown_col).value
            withdrawal = ws.cell(row=2, column=withdrawal_col).value
            self.assertGreater(tax, 0.0)
            self.assertAlmostEqual(withdrawal, drawdown + tax, places=6)
        finally:
            db.clear_section("assets")
            db.clear_section("expenses")

        # And with the assumption turned off, the Tax column goes to
        # zero and Withdrawal collapses back to Yearly Drawdown.
        db.add_line_item("assets", {"kind": "asset", "category": "401k", "name": "Fidelity 401k", "balance": 1_000_000, "include_in_monte_carlo": 1})
        db.add_line_item("expenses", {"name": "Baseline", "annual_amount": 100_000, "start_year": None, "end_year": None})
        db.set_assumptions({"apply_drawdown_tax": "0"})
        try:
            path = excel_export.export_current_state()
            wb = openpyxl.load_workbook(str(path))
            ws = wb["Monte Carlo Stress Test"]
            headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
            tax_col = headers.index("Great Depression (1929) Tax") + 1
            withdrawal_col = headers.index("Great Depression (1929) Withdrawal") + 1
            drawdown_col = headers.index("Yearly Drawdown") + 1
            tax = ws.cell(row=2, column=tax_col).value
            drawdown = ws.cell(row=2, column=drawdown_col).value
            withdrawal = ws.cell(row=2, column=withdrawal_col).value
            self.assertEqual(tax, 0.0)
            self.assertEqual(withdrawal, drawdown)
        finally:
            db.clear_section("assets")
            db.clear_section("expenses")
            db.set_assumptions({"apply_drawdown_tax": "1"})

    def test_stress_test_sheet_legacy_shared_columns_when_pools_not_supplied(self):
        # A caller that still hands in its own pools=None result (an
        # older test, or a future non-pools caller) keeps getting the
        # original shared Tax/Total Withdrawal layout -- this program
        # never drops support for that shape, only stops using it by
        # default.
        from source.reporting.monte_carlo import run_historical_stress_test

        stress_result = run_historical_stress_test(
            starting_balance=1_000_000,
            expenses=[{"name": "Baseline", "annual_amount": 100_000, "start_year": None, "end_year": None}],
            income=[],
            assumptions=db.get_assumptions(),
        )
        self.assertFalse(stress_result["per_scenario_tax"])
        path = excel_export.export_current_state(monte_carlo_result={"stress_test": stress_result})
        wb = openpyxl.load_workbook(str(path))
        ws = wb["Monte Carlo Stress Test"]
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        self.assertEqual(
            headers[:6],
            ["Year", "Total Expense", "Total Income", "Yearly Drawdown", "Tax (Fed+CA)", "Total Withdrawal"],
        )
        self.assertIn("Great Depression (1929) Return", headers)
        self.assertIn("Great Depression (1929) Balance", headers)
        self.assertNotIn("Great Depression (1929) Tax", headers)

    def test_default_monte_carlo_result_passes_pools_matching_the_live_tab(self):
        # Regression test for the export-vs-tab divergence this was
        # fixed to close: _default_monte_carlo_result() must hand the
        # same retirement_pools split to the simulations that the Monte
        # Carlo tab itself uses (source.webapp.pages.monte_carlo_page.
        # _render_monte_carlo_page), not the legacy pools=None mode.
        from source.reporting.aggregate import build_summary

        summary = build_summary()
        result = excel_export._default_monte_carlo_result(summary)
        self.assertTrue(result["stress_test"]["per_scenario_tax"])
        self.assertTrue(result["gaussian"].get("pools"))
        self.assertTrue(result["sp500_bootstrap"].get("pools"))


class TestAssumptionCell(unittest.TestCase):
    def test_non_finite_text_is_passed_through_instead_of_raising(self):
        for text in ("nan", "inf", "-inf"):
            self.assertEqual(excel_export._assumption_cell("expected_return", text), text)

    def test_ordinary_numbers_still_format(self):
        self.assertEqual(excel_export._assumption_cell("plan_horizon_years", "40").value, 40)


class TestExcelExportPresentation(TempConfigTestCase):
    """The export is meant to be opened and skimmed (or handed to someone),
    so it's styled consistently and carries native Excel charts."""

    def setUp(self):
        super().setUp()
        db.add_line_item("assets", {"kind": "asset", "category": "401k", "name": "Fidelity", "institution": "Fidelity", "balance": 100_000, "include_in_monte_carlo": 1})
        db.add_line_item("assets", {"kind": "asset", "category": "Cash", "name": "Checking", "institution": "Chase", "balance": 20_000})
        db.add_line_item("expenses", {"name": "Baseline", "category": "Fixed", "annual_amount": 40_000, "start_year": 2030})
        db.add_line_item("income", {"name": "Pension", "category": "Benefit", "annual_amount": 12_000})
        db.add_line_item("real_estate", {"name": "Home", "current_value": 500_000, "loan_balance": 200_000})
        for day, value in (("2026-09-01", 1_000_000.0), ("2026-09-02", 1_020_000.0), ("2026-09-03", 1_050_000.0)):
            db.record_snapshot(
                net_worth=value, manual_asset_total=value, linked_asset_total=0.0, liability_total=0.0,
                real_estate_equity_total=0.0, retirement_liquid_total=value, snapshot_date=day,
            )
        self.wb = openpyxl.load_workbook(str(excel_export.export_current_state()))

    @staticmethod
    def _chart_count(ws) -> int:
        return len(ws._charts)

    def test_every_data_sheet_has_a_styled_frozen_header(self):
        for name in ("Asset", "Expense", "Income", "Real Estate Analysis", "Net Worth History"):
            ws = self.wb[name]
            self.assertEqual(ws.freeze_panes, "A2", name)
            header = ws["A1"]
            self.assertTrue(header.font.bold, name)
            self.assertEqual(header.font.name, "Arial", name)
            self.assertEqual(header.fill.fgColor.rgb[-6:], "1D1D1F", name)

    def test_money_percent_and_year_cells_use_real_number_formats(self):
        ws = self.wb["Asset"]
        balance = next(ws.cell(row=r, column=6) for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=3).value == "Fidelity")
        self.assertIn("$", balance.number_format)
        ws = self.wb["Expense"]
        self.assertEqual(ws["D2"].number_format, "0")  # Start Year is a year, not 2,026.00 money

    def test_summary_sheet_has_charts_and_highlights_the_key_figures(self):
        ws = self.wb["Summary"]
        self.assertGreaterEqual(self._chart_count(ws), 2)
        labels = {ws.cell(row=r, column=1).value: r for r in range(1, ws.max_row + 1)}
        net_worth_row = labels["Net worth"]
        self.assertTrue(ws.cell(row=net_worth_row, column=1).font.bold)

    def test_net_worth_history_sheet_lists_snapshots_and_plots_them(self):
        ws = self.wb["Net Worth History"]
        dates = [ws.cell(row=r, column=1).value for r in range(2, 5)]
        self.assertEqual(dates, ["2026-09-01", "2026-09-02", "2026-09-03"])
        self.assertEqual(ws["B4"].value, 1_050_000.0)
        self.assertEqual(self._chart_count(ws), 1)

    def test_chart_series_start_at_the_first_data_row_not_the_header(self):
        # A series reference that began on the header row would plot the
        # header text as a zero-valued first point.
        ws = self.wb["Net Worth History"]
        chart = ws._charts[0]
        ref = chart.series[0].val.numRef.f
        self.assertIn("$B$2", ref)
        self.assertNotIn("$B$1", ref)

    def test_net_worth_history_with_a_single_snapshot_has_a_note_instead_of_a_chart(self):
        db.clear_net_worth_snapshots() if hasattr(db, "clear_net_worth_snapshots") else None
        from source.common.db.connection import connect

        with connect() as conn:
            conn.execute("DELETE FROM net_worth_snapshots WHERE snapshot_date != '2026-09-01'")
        wb = openpyxl.load_workbook(str(excel_export.export_current_state()))
        ws = wb["Net Worth History"]
        self.assertEqual(self._chart_count(ws), 0)
        values = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
        self.assertTrue(any(isinstance(v, str) and "captured automatically" in v for v in values))

    def test_monte_carlo_sheet_has_fan_charts_with_percentage_success_rate(self):
        ws = self.wb["Monte Carlo Simulation"]
        self.assertEqual(self._chart_count(ws), 2)  # Gaussian + S&P 500
        for chart in ws._charts:
            self.assertEqual(len(chart.series), 3)  # 90th, median, 10th
        row = next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value == "Success rate")
        self.assertIn("%", ws.cell(row=row, column=2).number_format)
        self.assertLessEqual(ws.cell(row=row, column=2).value, 1.0)  # stored as a fraction

    def test_stress_test_sheet_charts_every_scenario_balance(self):
        ws = self.wb["Monte Carlo Stress Test"]
        self.assertEqual(self._chart_count(ws), 2)
        balance_chart = ws._charts[0]
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        scenarios = sum(1 for h in headers if isinstance(h, str) and h.endswith(" Balance"))
        self.assertEqual(len(balance_chart.series), scenarios)

    def test_asset_expense_income_and_real_estate_sheets_chart_their_breakdowns(self):
        for name in ("Asset", "Expense", "Income", "Real Estate Analysis"):
            self.assertGreaterEqual(self._chart_count(self.wb[name]), 1, name)

    def test_assumptions_are_written_as_numbers_with_readable_labels(self):
        ws = self.wb["Monte Carlo Simulation"]
        rows = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2) for r in range(1, 40)}
        cell = rows["Expected return"]
        self.assertIsInstance(cell.value, float)
        self.assertIn("%", cell.number_format)

    def test_sheets_have_print_setup(self):
        ws = self.wb["Asset"]
        self.assertEqual(ws.page_setup.orientation, "landscape")
        self.assertTrue(ws.sheet_properties.pageSetUpPr.fitToPage)

    def test_top_n_with_other_folds_the_tail_and_drops_non_positive_values(self):
        from source.reporting.excel_style import top_n_with_other

        items = [(f"c{i}", 100.0 - i) for i in range(10)] + [("neg", -5.0), ("zero", 0.0)]
        folded = top_n_with_other(items, n=6)
        self.assertEqual(len(folded), 7)
        self.assertEqual(folded[-1][0], "Other")
        self.assertAlmostEqual(folded[-1][1], sum(100.0 - i for i in range(6, 10)))
        self.assertEqual([k for k, _ in folded[:3]], ["c0", "c1", "c2"])

    def test_fitted_axis_min_leaves_headroom_below_the_data(self):
        from source.reporting.excel_style import fitted_axis_min

        floor = fitted_axis_min([7_100_000.0, 8_050_000.0])
        self.assertLess(floor, 7_100_000.0)
        self.assertGreater(floor, 6_000_000.0)
        self.assertIsNone(fitted_axis_min([]))
        self.assertIsNone(fitted_axis_min([-5.0, 10.0]))


if __name__ == "__main__":
    unittest.main()
